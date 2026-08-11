import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_pdf_text(file_path: str) -> str:
    try:
        import fitz
        doc = fitz.open(file_path)
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                pages.append(f"--- Page {i + 1} ---\n{text.strip()}")
        doc.close()
        full_text = "\n\n".join(pages)
        if not full_text.strip():
            return "[PDF contains no extractable text — may be image-based]"
        max_chars = 8000
        if len(full_text) > max_chars:
            full_text = full_text[:max_chars] + "\n\n[... truncated ...]"
        return full_text
    except Exception as e:
        logger.error("PDF extraction failed: %s", e)
        return f"[PDF extraction error: {e}]"


def extract_docx_text(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            return "[DOCX contains no text content]"
        max_chars = 8000
        if len(full_text) > max_chars:
            full_text = full_text[:max_chars] + "\n\n[... truncated ...]"
        return full_text
    except Exception as e:
        logger.error("DOCX extraction failed: %s", e)
        return f"[DOCX extraction error: {e}]"


def extract_document(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(file_path)
    elif ext in (".docx", ".doc"):
        return extract_docx_text(file_path)
    else:
        return f"[Unsupported document format: {ext}]"
